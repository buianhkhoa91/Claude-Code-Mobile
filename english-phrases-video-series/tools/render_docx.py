# -*- coding: utf-8 -*-
"""Render 300 video-short scripts (.docx) + table of contents from lesson JSON data.

Usage: python3 render_docx.py
Reads:  data/outline.json, data/lessons-*.json
Writes: 00-Muc-Luc-300-Bai.docx, scripts-docx/Bai-NNN-<slug>.docx
"""
import glob
import json
import os
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, "data")
OUT = os.path.join(BASE, "scripts-docx")

SERIES_TITLE = "300 Bài Cụm Từ Tiếng Anh Thông Dụng Cho Người Đi Làm"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x59, 0x59, 0x59)
GREEN = RGBColor(0x2E, 0x7D, 0x32)


def slugify(text):
    text = unicodedata.normalize("NFD", text)
    text = text.replace("đ", "d").replace("Đ", "D")
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-")
    return text[:60].rstrip("-")


def load_lessons():
    lessons = {}
    for path in sorted(glob.glob(os.path.join(DATA, "lessons-*.json"))):
        for item in json.load(open(path, encoding="utf-8")):
            lessons[item["lesson"]] = item
    return lessons


def add_heading(doc, text, size=14, color=ACCENT, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def add_line(doc, label, text, label_color=GRAY, text_bold=False, text_color=None, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    if label:
        run = p.add_run(label + " ")
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = label_color
    run = p.add_run(text)
    run.bold = text_bold
    run.font.size = Pt(size)
    if text_color:
        run.font.color.rgb = text_color
    return p


def render_lesson(lesson, next_lesson_label):
    n = lesson["lesson"]
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BÀI %d – %s" % (n, lesson["topic"]))
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(SERIES_TITLE + "  |  Video short ~60 giây  |  12 Scene")
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY

    # Scene 1: Hook
    add_heading(doc, "SCENE 1 – HOOK (MỞ ĐẦU)")
    add_line(doc, "Hiển thị:", "Bài %d – %s" % (n, lesson["topic"]))
    add_line(doc, "Voice (tiếng Việt):", lesson["hook"])

    # Scenes 2-11: 5 phrases x 2 scenes
    scene = 2
    for i, ph in enumerate(lesson["phrases"], 1):
        add_heading(doc, "SCENE %d – CỤM TỪ %d" % (scene, i))
        add_line(doc, "Voice (Anh-Mỹ, đọc):", ph["en"], text_bold=True, text_color=GREEN, size=12)
        add_line(doc, "IPA:", ph["ipa"])
        add_line(doc, "Nghĩa (chỉ hiển thị, không đọc):", ph["vi"])
        scene += 1

        add_heading(doc, "SCENE %d – VÍ DỤ CỤM TỪ %d" % (scene, i))
        add_line(doc, "Voice (Anh-Mỹ, đọc):", ph["example"], text_bold=True, text_color=GREEN, size=12)
        add_line(doc, "Dịch (chỉ hiển thị, không đọc):", ph["example_vi"])
        scene += 1

    # Scene 12: CTA
    add_heading(doc, "SCENE 12 – KẾT THÚC (CTA)")
    add_line(doc, "Voice (tiếng Việt):",
             "Nghe và lặp lại để ghi nhớ bạn nhé! Đừng quên theo dõi bài học tiếp theo: %s." % next_lesson_label)
    add_line(doc, "Hiển thị:", "Nghe & lặp lại  •  Theo dõi: %s" % next_lesson_label)

    fname = "Bai-%03d-%s.docx" % (n, slugify(lesson["topic"]))
    doc.save(os.path.join(OUT, fname))
    return fname


def render_toc(outline, lessons):
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MỤC LỤC – " + SERIES_TITLE)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = ACCENT

    intro = doc.add_paragraph()
    run = intro.add_run(
        "300 video short (~60 giây/video) cho người Việt đi làm. Mỗi bài gồm 5 cụm từ "
        "(phrasal verbs, idioms, collocations) kèm IPA chuẩn Oxford (giọng Anh-Mỹ), nghĩa tiếng Việt, "
        "câu ví dụ và bản dịch. Cấu trúc video: 1 Scene Hook + 10 Scene nội dung (mỗi cụm từ 2 scene) + 1 Scene CTA.")
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY

    for module in outline["modules"]:
        add_heading(doc, "CHƯƠNG %d: %s (Bài %d–%d)" % (
            module["module"], module["title"],
            module["lessons"][0]["lesson"], module["lessons"][-1]["lesson"]), size=13)
        for item in module["lessons"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run("Bài %d: %s" % (item["lesson"], item["topic"]))
            run.font.size = Pt(11)
            data = lessons.get(item["lesson"])
            if data:
                run = p.add_run("  —  " + " | ".join(ph["en"] for ph in data["phrases"]))
                run.font.size = Pt(9)
                run.font.color.rgb = GRAY

    doc.save(os.path.join(BASE, "00-Muc-Luc-300-Bai.docx"))


def main():
    os.makedirs(OUT, exist_ok=True)
    outline = json.load(open(os.path.join(DATA, "outline.json"), encoding="utf-8"))
    topics = {l["lesson"]: l["topic"] for m in outline["modules"] for l in m["lessons"]}
    lessons = load_lessons()

    missing = [n for n in range(1, 301) if n not in lessons]
    if missing:
        print("MISSING lessons:", missing)

    count = 0
    for n in sorted(lessons):
        if n < 300:
            next_label = "Bài %d – %s" % (n + 1, topics[n + 1])
        else:
            next_label = "Bài 1 – %s (ôn lại từ đầu chuỗi 300 bài)" % topics[1]
        render_lesson(lessons[n], next_label)
        count += 1

    render_toc(outline, lessons)
    print("Rendered %d lesson docx files + 00-Muc-Luc-300-Bai.docx" % count)


if __name__ == "__main__":
    main()
